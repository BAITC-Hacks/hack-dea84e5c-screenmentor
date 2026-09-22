from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "docx" / "architecture_and_privacy.docx"
LOGO = ROOT / "extension" / "mentor-logo.png"
GREEN = "205944"
PALE = "EAF4F0"
GRAY = "F5F7F6"
GRID = "D9D9D9"
TEXT = RGBColor(30, 47, 40)

def font(run, size=10.4, bold=False, color=TEXT):
    run.font.name = "Arial"; run._element.rPr.rFonts.set(qn("w:ascii"), "Arial"); run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size); run.bold = bold; run.font.color.rgb = color

def text(p, value, **kw):
    r = p.add_run(value); font(r, **kw); return r

def shade(cell, value):
    props = cell._tc.get_or_add_tcPr(); node = props.find(qn("w:shd"))
    if node is None: node = OxmlElement("w:shd"); props.append(node)
    node.set(qn("w:fill"), value)

def border(cell):
    props = cell._tc.get_or_add_tcPr(); node = props.first_child_found_in("w:tcBorders")
    if node is None: node = OxmlElement("w:tcBorders"); props.append(node)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = node.find(qn(f"w:{edge}"))
        if e is None: e = OxmlElement(f"w:{edge}"); node.append(e)
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "8"); e.set(qn("w:color"), GRID)

def margin(cell):
    props = cell._tc.get_or_add_tcPr(); node = props.first_child_found_in("w:tcMar")
    if node is None: node = OxmlElement("w:tcMar"); props.append(node)
    for side in ("top", "start", "bottom", "end"):
        e = node.find(qn(f"w:{side}"))
        if e is None: e = OxmlElement(f"w:{side}"); node.append(e)
        e.set(qn("w:w"), "120"); e.set(qn("w:type"), "dxa")

def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"; t.autofit = False
    for i, value in enumerate(headers):
        c = t.rows[0].cells[i]; c.width = Cm(widths[i]); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; shade(c, GREEN); border(c); margin(c)
        p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); text(p, value, size=9.1, bold=True, color=RGBColor(255,255,255))
    trp = t.rows[0]._tr.get_or_add_trPr(); h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trp.append(h)
    for n, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            c = cells[i]; c.width = Cm(widths[i]); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; shade(c, "FFFFFF" if n % 2 == 0 else GRAY); border(c); margin(c)
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.06; text(p, str(value), size=8.8)
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)

def heading(doc, value, level=1):
    p = doc.add_paragraph(style=f"Heading {level}"); p.paragraph_format.keep_with_next = True; p.paragraph_format.space_before = Pt(14 if level == 1 else 9); p.paragraph_format.space_after = Pt(5)
    text(p, value, size=15 if level == 1 else 11.8, bold=True, color=RGBColor(0,0,0))

def body(doc, value, lead=None):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(7); p.paragraph_format.line_spacing = 1.16
    if lead: text(p, lead, bold=True)
    text(p, value)

def bullet(doc, value):
    p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(3); p.paragraph_format.line_spacing = 1.08; text(p, value)

def setup(doc):
    s = doc.sections[0]; s.top_margin=Cm(1.65); s.bottom_margin=Cm(1.55); s.left_margin=Cm(1.75); s.right_margin=Cm(1.75)
    f = s.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER; f.paragraph_format.space_before=Pt(6); text(f, "ИИ Наставник  |  Архитектура и конфиденциальность  |  Версия 0.24.20", size=8.3, color=RGBColor(100,112,106))
    normal=doc.styles["Normal"]; normal.font.name="Arial"; normal._element.rPr.rFonts.set(qn("w:ascii"),"Arial"); normal._element.rPr.rFonts.set(qn("w:hAnsi"),"Arial"); normal.font.size=Pt(10.4)
    for name in ("Title", "Heading 1", "Heading 2"):
        st=doc.styles[name]; st.font.name="Arial"; st._element.rPr.rFonts.set(qn("w:ascii"),"Arial"); st._element.rPr.rFonts.set(qn("w:hAnsi"),"Arial"); st.font.color.rgb=RGBColor(0,0,0)
    ppr=doc.styles["Title"]._element.get_or_add_pPr()
    for b in ppr.findall(qn("w:pBdr")): ppr.remove(b)

def build():
    doc=Document(); setup(doc)
    if LOGO.exists():
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(7); p.add_run().add_picture(str(LOGO), width=Cm(2.25))
    p=doc.add_paragraph(style="Title"); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(7); text(p,"Архитектура и конфиденциальность ИИ Наставника",size=23,bold=True,color=RGBColor(0,0,0))
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(18); text(p,"Технический обзор для Demo Day и пилотных партнёров",size=11.5,color=RGBColor(75,95,86))
    body(doc,"ИИ Наставник - браузерное расширение, которое объясняет элементы интерфейса и предлагает один безопасный следующий шаг. Архитектура построена вокруг принципа минимизации данных: помощник не совершает действия за человека и получает только ограниченные описатели видимых элементов, необходимые для ответа.",lead="Главный вывод. ")
    table(doc,["Версия","Назначение","Ключевая гарантия"],[["0.24.20","Демонстрация, пилоты, техническая проверка","Пользователь сохраняет контроль: никаких автоматических кликов, отправки форм или платежей"]],[3.0,5.0,9.0])
    body(doc,"Документ описывает реализованную схему и её ограничения. Она снижает объём обрабатываемых данных, но не является обещанием абсолютной анонимности: на обычном сайте безопасная подпись элемента теоретически может содержать персональную информацию. Для реальных пилотов нужны учебные данные, отдельная privacy-проверка и согласованные домены.",lead="Граница обещания. ")
    doc.add_page_break()

    heading(doc,"1  Роли компонентов")
    table(doc,["Компонент","Ответственность","Что не делает"],[
        ["Пользователь и сайт","Открывает панель, задаёт текстовый или голосовой вопрос, выбирает видимый элемент и сам подтверждает действие.","Не передаёт расширению право действовать от своего имени."],
        ["Chrome расширение","Показывает панель в Shadow DOM, локально фильтрует интерфейсные признаки, хранит краткую сессию вкладки и управляет записью.","Не читает полный HTML, текст документа, содержимое форм, cookies, history, clipboardRead или debugger."],
        ["Node.js backend","Проверяет pairing token и входные схемы, ограничивает запросы, отменяет устаревшие вызовы, держит короткий in-memory cache.","Не хранит постоянную историю задач, ответы или полный DOM."],
        ["Внешние сервисы","LLM формирует объяснение; Azure Speech обрабатывает явно запущенный голос; Google Safe Browsing проверяет хеш-префиксы URL.","Не получают ключей внутри пакета расширения или необработанный DOM."],
    ],[3.2,7.4,6.4])
    heading(doc,"Принцип работы",level=2)
    for item in [
        "Модель предлагает только один проверяемый следующий шаг или просит открыть нужный раздел. Она не получает право нажимать кнопки, менять данные, отправлять формы или проводить платежи.",
        "Для Word и Excel используются подготовленные каталоги команд. На остальных сайтах назначение элемента определяется моделью по текущим очищенным признакам, без словарей исключений по конкретным сайтам.",
        "Ответ проходит строгую схему: backend отклоняет ответ с несуществующей целью, неверным форматом или неподходящим языком и может выполнить безопасную повторную генерацию.",
    ]: bullet(doc,item)
    heading(doc,"2  Поток запроса",level=1)
    table(doc,["Шаг","Что происходит","Защитный контроль"],[
        ["Наблюдение","Пользователь выбирает видимый элемент или нажимает «Изучить страницу». Расширение строит краткий снимок интерфейсных признаков.","Локальный фильтр исключает известные личные области, чувствительные поля, адреса и характерные личные маркеры."],
        ["Подготовка","В запрос включаются язык, задача, UUID и ограниченный набор descriptors: тип, роль, очищенная подпись, состояние и структура.","До 18 описателей для обучения и до 36 для вопроса; полный DOM и снимки экрана не уходят."],
        ["Проверка","Backend проверяет источник, token, лимиты и JSON-схему. Запрос может быть отменён при изменении страницы или новой задаче.","Bearer pairing, allowlist схем, rate limits, отсечение поздних ответов."],
        ["Ответ","Пользователь получает объяснение и подсветку. Если переход нужен, помощник называет его, но пользователь нажимает сам.","Подсветка не перекрывает выбранный элемент; нет программного клика от имени агента."],
    ],[2.6,8.2,6.2])
    doc.add_page_break()

    heading(doc,"3  Какие данные обрабатываются")
    table(doc,["Категория","Разрешено для функции","Не читается и не отправляется"],[
        ["Контекст интерфейса","Тип control: button, link, input, select, checkbox; роль, очищенная доступная подпись, disabled/state и ограниченная структура.","Полный HTML, скрипты, стили, class/id, содержимое редактируемого документа и скриншоты."],
        ["Содержимое пользователя","Только текст вопроса, который пользователь сам вводит или подтверждает голосом.","Текст документов, ячеек Excel, чатов, письмо, содержимое форм, пароли, OTP и платёжные данные."],
        ["Локальное состояние","Язык, внешний вид, pairing token, прогресс урока Office и краткая сессия вкладки.","История посещений, cookies, identity, постоянный журнал наблюдаемых элементов."],
        ["Диагностика","UUID, язык, тип задачи, число элементов, длительность, cache hit и категория ошибки.","Домены, подписи, ответы модели, ключи API и секреты."],
    ],[3.1,7.2,6.7])
    body(doc,"Фильтры эвристические. На произвольном сайте имя человека может оказаться в обычной кнопке, поэтому продукт не заявляет, что каждая подпись всегда является публичной. Для демонстрации и обучения рекомендуется использовать синтетические данные.",lead="Практическое ограничение. ")
    heading(doc,"4  Хранение и жизненный цикл")
    table(doc,["Где","Что хранится","Срок и защита"],[
        ["Память вкладки","Временные DOM-ссылки и текущий снимок до 36 разрешённых элементов.","Только память текущей вкладки; не записывается на диск."],
        ["Хранилище расширения","Язык, прогресс Office, внешний вид и pairing token.","Token живёт в IndexedDB origin расширения; popup может удалить его, но не показывает обратно. Это не системное хранилище паролей."],
        ["Backend memory cache","До 64 безопасных подсказок: домен, разрешённые признаки и ответ.","До 10 минут; перезапуск backend очищает cache."],
        ["Backend secrets","Ключи провайдеров и серверные настройки.","Только в .env.local на backend; ключи не входят в пакет расширения и не попадают в логи."],
    ],[3.0,7.4,6.6])
    doc.add_page_break()

    heading(doc,"5  Отдельные маршруты данных")
    table(doc,["Функция","Маршрут","Контроль и границы"],[
        ["Текстовый вопрос","Расширение -> backend -> LLM -> backend -> панель.","Передаются только вопрос и краткие разрешённые descriptors; ответ проверяется схемой и языком."],
        ["Голосовой вопрос","Микрофон -> Azure Speech to Text -> текст -> backend -> LLM.","Запись запускает сам пользователь. WAV 16 kHz mono ограничен 55 сек.; аудио не хранится после распознавания."],
        ["Озвучивание ответа","Текст видимого ответа -> speechSynthesis браузера.","Это клиентская функция; при отсутствии подходящего голоса показывается понятная ошибка."],
        ["Проверка риска сайта","URL текущей вкладки -> backend -> SHA-256 hash-prefix -> Google Safe Browsing.","Google получает только уникальные 4-байтовые префиксы хешей, а не URL. Полное совпадение проверяется локально на backend."],
    ],[3.1,7.4,6.5])
    heading(doc,"Проверка риска сайта",level=2)
    body(doc,"Расширение запрашивает проверку при открытии панели наставника. Чистая страница не вызывает сообщение. При совпадении категорий MALWARE, SOCIAL_ENGINEERING или UNWANTED_SOFTWARE, либо при локальных признаках адреса, появляется небольшое закрываемое предупреждение рядом с иконкой агента. Оно не блокирует сайт, не закрывает вкладку и не является антивирусной гарантией.")
    for item in [
        "Полный URL нужен backend только для нормализации и вычисления хешей. При удалённом backend адрес доходит до сервера по HTTPS, поэтому production-сервер должен контролироваться командой.",
        "Ответ Google V5 декодируется локально только для полного хеша, категории угрозы и срока cache. URL и хеши не записываются в логи.",
        "Локальные признаки - HTTP, числовой IP, нестандартный порт, международный домен или сложное имя - означают «нужна осторожность», но не доказывают мошенничество.",
    ]: bullet(doc,item)
    heading(doc,"6  Контроль доступа и эксплуатация")
    table(doc,["Контроль","Реализация"],[
        ["Разрешения браузера","Доступ к HTTP/HTTPS страницам нужен для панели. Разрешения cookies, history, identity, debugger, downloads и clipboardRead не запрашиваются."],
        ["Подключение к backend","Pairing code создаёт bearer token. Управляющие сообщения принимает только popup того же расширения; «Забыть подключение» удаляет token и адрес backend."],
        ["Секреты",".env.local не включается в git и релизный ZIP. Диагностика показывает только configured/missing, без печати ключа."],
        ["Надёжность","Health endpoint прогревается перед демонстрацией; отмена задач предотвращает показ устаревшего ответа, но уже обработанные токены у провайдера могут учитываться по его условиям."],
    ],[4.2,12.8])
    heading(doc,"Целевая модель доступа для production",level=2)
    body(doc,"Текущий pairing-token остаётся только для локальной разработки, хакатона и технического стенда. Для обычного пользователя планируется вход по номеру телефона или email с одноразовым кодом либо ограниченный режим без регистрации с небольшой квотой. Пользователь не должен вручную вводить LOCAL_ACCESS_TOKEN, адрес backend или ключ модели.")
    body(doc,"Школы и другие учебные организации смогут выдавать код организации, invite-ссылку или QR-код с ограниченным сроком действия, числом активаций и отдельной квотой. Подтверждение статуса школьника, студента или пожилого пользователя рассматривается только через официальный eGov/ЭЦП-поток после отдельной юридической и privacy-проверки. ИИН не является паролем, не отправляется модели и не хранится в расширении.")
    doc.add_page_break()

    heading(doc,"7  Проверяемые ограничения и план пилота")
    body(doc,"Ниже - условия, которые важно проверить до внешнего пилота. Они переводят privacy-by-design из заявления в воспроизводимый набор проверок.",lead="Цель пилота. ")
    table(doc,["Проверка","Критерий прохождения"],[
        ["Payload inspection","Тестовые маркеры пароля, OTP, текста письма и ячейки не появляются в payload, TTS или production-логах."],
        ["UI safety","Подсказка и coachmark не перекрывают целевой элемент; нажатие всегда остаётся действием пользователя."],
        ["Language and persistence","Новая установка открывается на казахском; выбор языка вручную сохраняется отдельно для панели и popup."],
        ["Voice consent","Браузерное разрешение на микрофон запрашивается только в пользовательском действии; остановка записи отправляет текст в существующий сценарий вопроса."],
        ["Safe Browsing","Чистый URL не показывает notice; тестовое совпадение показывает предупреждение без передачи URL в Google."],
        ["Deployment","HTTPS backend, актуальные секреты, health check и ручной demo warm-up проверены на чистом профиле браузера."],
    ],[4.5,12.5])
    heading(doc,"Заключение",level=1)
    body(doc,"Архитектура ИИ Наставника намеренно ограничивает возможности агента: он объясняет, но не действует; передаёт минимум интерфейсного контекста, а не страницу целиком; хранит краткое техническое состояние, а не историю пользователя. Такой подход делает продукт понятнее для начинающих пользователей и даёт команде проверяемую основу для пилотов.")
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(8); p.paragraph_format.space_after=Pt(0); text(p,"Следующий шаг: ",bold=True); text(p,"провести пилот на учебных данных, измерить долю успешно завершённых сценариев, число повторных вопросов и понятность предупреждений о риске.")
    OUT.parent.mkdir(parents=True,exist_ok=True); doc.core_properties.title="Архитектура и конфиденциальность ИИ Наставника"; doc.core_properties.author="HackAlem AI"; doc.core_properties.subject="Технический обзор"; doc.save(OUT)

if __name__ == "__main__": build()
