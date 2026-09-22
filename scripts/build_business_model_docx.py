from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "docx" / "ai_mentor_business_model.docx"
LOGO = ROOT / "extension" / "mentor-logo.png"

TEAL = "205944"
PALE_TEAL = "EAF4F0"
PALE_GRAY = "F5F7F6"
GRID = "D9D9D9"
TEXT = RGBColor(30, 47, 40)


def set_cell_shading(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_border(cell, color=GRID, size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_cell_margins(cell, top=110, start=125, bottom=110, end=125):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    element = OxmlElement("w:tblHeader")
    element.set(qn("w:val"), "true")
    tr_pr.append(element)


def set_keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def set_font(run, size=10.5, bold=False, color=TEXT):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run._element.rPr.rFonts.set(qn("w:cs"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def add_text(paragraph, text, bold=False, size=10.5, color=TEXT):
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=bold, color=color)
    return run


def add_body(doc, text, lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.17
    if lead:
        add_text(p, lead, bold=True)
    add_text(p, text)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.1
    add_text(p, text)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    add_text(p, text, size=15 if level == 1 else 12, bold=True, color=RGBColor(0, 0, 0))
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.style = "Table Grid"
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        if widths:
            cell.width = Cm(widths[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, TEAL)
        set_cell_border(cell)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_text(p, header, bold=True, size=9.3, color=RGBColor(255, 255, 255))
    set_repeat_table_header(table.rows[0])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for col_index, value in enumerate(row):
            cell = cells[col_index]
            if widths:
                cell.width = Cm(widths[col_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cell, PALE_GRAY if row_index % 2 else "FFFFFF")
            set_cell_border(cell)
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            add_text(p, str(value), size=8.7)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return table


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    add_text(p, "ИИ Наставник  |  Бизнес модель  |  Версия 0.24.20", size=8.5, color=RGBColor(100, 112, 106))


def set_document_defaults(doc):
    section = doc.sections[0]
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.55)
    section.left_margin = Cm(1.75)
    section.right_margin = Cm(1.75)
    add_footer(section)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT
    for name in ("Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.color.rgb = RGBColor(0, 0, 0)
    title_pr = doc.styles["Title"]._element.get_or_add_pPr()
    for border in title_pr.findall(qn("w:pBdr")):
        title_pr.remove(border)


def build():
    doc = Document()
    set_document_defaults(doc)

    if LOGO.exists():
        logo = doc.add_paragraph()
        logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        logo.paragraph_format.space_after = Pt(8)
        logo.add_run().add_picture(str(LOGO), width=Cm(2.25))

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(7)
    add_text(title, "Бизнес модель ИИ Наставника", size=25, bold=True, color=RGBColor(0, 0, 0))
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(20)
    add_text(subtitle, "Для Demo Day и пилотных партнёров", size=12, color=RGBColor(75, 95, 86))

    add_body(doc, "Документ описывает, как превратить браузерного помощника для русско- и казахоязычных пользователей в проверяемый B2B2C-продукт. Основная рекомендация: начать с контролируемых пилотов в школах, библиотеках и центрах цифровой грамотности, где можно измерить пользу без доступа к чувствительным данным.", lead="Главный вывод. ")
    meta = add_table(doc, ["Версия", "Статус", "Цель документа"], [["0.24.20", "Рабочая гипотеза", "Пилотные переговоры, Demo Day и план первых продаж"]], [3.0, 4.0, 10.0])
    doc.add_paragraph()
    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(5)
    note.paragraph_format.space_after = Pt(0)
    add_text(note, "Важно. ", bold=True, size=9.5)
    add_text(note, "Цены, конверсии и объём рынка в документе являются гипотезами до интервью и пилотных измерений.", size=9.5, color=RGBColor(78, 91, 84))
    doc.add_page_break()

    add_heading(doc, "1  Продукт и проблема")
    add_body(doc, "ИИ Наставник - расширение для Chrome и Edge. Оно объясняет видимые элементы текущего веб-интерфейса простым языком, принимает текстовый и голосовой вопрос, озвучивает ответ и показывает предупреждение при признаках рискованного сайта.")
    add_body(doc, "Продукт не действует вместо человека: не нажимает кнопки, не заполняет формы, не отправляет сообщения и не проводит платежи. Это снижает риск ошибочной автоматизации и сохраняет контроль у пользователя.")
    add_heading(doc, "Что меняется для пользователя", level=2)
    for item in [
        "Вместо поиска общей инструкции человек получает пояснение рядом с текущей кнопкой, полем или разделом.",
        "Вместо догадки о следующем действии он получает один безопасный шаг или честную просьбу открыть нужное меню.",
        "Вместо языкового барьера доступны казахский и русский языки, голосовой ввод и озвучивание.",
        "Вместо ложного обещания полной защиты показывается аккуратное предупреждение только при признаках риска.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "2  Приоритетные сегменты")
    add_table(doc,
        ["Сегмент", "Кто платит", "Боль", "Первая ценность"],
        [
            ["Пожилые и начинающие", "Семья, социальный или образовательный партнёр", "Страх ошибиться и непонимание интерфейса", "Понятная подсказка, голос, два языка"],
            ["Школьники и студенты", "Школа, колледж, вуз, EdTech", "Не знает, как выполнить задачу в веб-сервисе", "Обучение интерфейсу и Word/Excel без чтения файла"],
            ["Библиотеки и НКО", "Организация", "Повторяющиеся простые вопросы и очереди", "Первый уровень цифровой помощи"],
            ["Корпоративный onboarding", "Компания или владелец сервиса", "Долгое освоение систем и однотипные обращения", "Помощь на разрешённых страницах"],
        ], [3.25, 3.65, 4.9, 5.2])
    add_body(doc, "Стартовый фокус - B2B2C-пилоты. Один партнёр даёт доступ к группе пользователей, помогает провести обучение и позволяет сравнить результат до и после использования продукта.", lead="Приоритет запуска. ")
    doc.add_page_break()

    add_heading(doc, "3  Ценность и продуктовое предложение")
    add_table(doc, ["Сторона", "Что получает"], [
        ["Пользователь", "Понятное объяснение на привычном языке, голосовой вопрос и безопасный следующий шаг без автоматического действия."],
        ["Семья", "Меньше тревоги при освоении сайтов близким человеком и видимый сигнал о подозрительных признаках адреса."],
        ["Преподаватель или консультант", "Меньше повторяющихся вопросов о базовой навигации и материал для самостоятельной практики."],
        ["Организация", "Более мягкий onboarding, потенциальное снижение однотипных обращений и инклюзивный цифровой сервис."],
    ], [4.2, 12.8])

    add_heading(doc, "Базовые возможности", level=2)
    for item in [
        "Плавающая панель в браузере рядом с интерфейсом.",
        "Объяснение видимых кнопок, полей, ссылок и структурных разделов без передачи полного HTML страницы.",
        "Вопрос своими словами: помощник связывает его только с наблюдаемой командой и не придумывает скрытый маршрут.",
        "Казахский язык по умолчанию для новой установки и сохранение выбранного языка.",
        "Голосовой вопрос после остановки записи сразу поступает в сценарий помощи; ответ можно озвучить.",
        "Предупреждение о риске сайта без навязчивого сообщения для чистых страниц.",
        "Подготовленные безопасные объяснения Word и Excel Online без чтения документа, ячеек или имени файла.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "Коммерческие пакеты  рабочая гипотеза", level=2)
    add_table(doc, ["Пакет", "Кому", "Содержание", "Коммерческая цель"], [
        ["Community pilot", "Школа, библиотека, НКО", "Ограниченная группа, базовые сценарии, сопровождение пилота", "Доказать пользу и собрать обратную связь"],
        ["Education", "Учебная организация", "Управляемое развёртывание, Word/Excel-сценарии, агрегированные метрики", "Годовая лицензия организации"],
        ["Team", "Компания или служба поддержки", "Согласованный перечень доменов, onboarding, приоритетная поддержка", "Подписка по активным пользователям или подразделениям"],
        ["Enterprise", "Крупный владелец сервиса", "Брендинг, SSO и аудит после отдельной проработки", "Внедрение и годовая поддержка"],
    ], [3.0, 3.2, 6.0, 4.8])
    add_heading(doc, "Модель доступа и регистрации", level=2)
    add_body(doc, "Обычный пользователь получает простой вход по номеру телефона или email с одноразовым кодом либо может начать работу без регистрации с небольшой квотой. Пользователю не показываются LOCAL_ACCESS_TOKEN, адрес backend и ключ модели.")
    add_body(doc, "Школы и другие учебные организации используют код организации, invite-ссылку или QR-код с ограниченным сроком действия, числом активаций, отзывом и отдельной квотой. Подтверждение статуса школьника, студента или пожилого пользователя рассматривается только через официальный eGov/ЭЦП-поток после privacy-проверки. ИИН не используется как пароль и не попадает в модель. Для HackAlem AI остаётся текущий pairing-токен как режим технического стенда.")
    doc.add_page_break()

    add_heading(doc, "4  Business Model Canvas")
    add_table(doc, ["Блок", "Решение для ИИ Наставника"], [
        ["Customer Segments", "Начинающие пользователи, пожилые, учащиеся, организации цифровой грамотности и команды поддержки."],
        ["Value Proposition", "Контекстная помощь рядом с интерфейсом на KK и RU; объясняет, но не совершает рискованное действие вместо человека."],
        ["Channels", "Пилоты с образовательными и социальными организациями, Chrome Web Store и Edge Add-ons, партнёрские демо, рекомендации преподавателей и семей."],
        ["Customer Relationships", "Простой onboarding, встроенная инструкция, база сценариев, поддержка партнёра и канал обратной связи."],
        ["Revenue Streams", "Организационная подписка, плата за внедрение и approved-domain профиль, корпоративная поддержка. B2C freemium - только после проверки удержания."],
        ["Key Activities", "Улучшение безопасных схем, двуязычного UX, доступности, эксплуатации backend и проведение пилотов."],
        ["Key Resources", "Расширение, backend, двуязычный UX, проверенный каталог Office-команд, методика тестирования и партнёрства."],
        ["Key Partners", "Школы, библиотеки, НКО, программы цифровой грамотности, AI и Speech-провайдеры, владельцы разрешённых сайтов."],
        ["Cost Structure", "Разработка, облачная инфраструктура, модель и Speech, безопасное хранение секретов, поддержка, исследование пользователей и privacy-проверка."],
    ], [4.6, 12.4])

    add_heading(doc, "5  Доходная модель и экономика")
    add_body(doc, "Пилоту не нужно продавать AI-токены. Партнёру понятнее оплачивать доступ организации, набор пользователей, сопровождение и согласованные сценарии. Команда при этом считает переменную стоимость на активного пользователя и не продаёт пакет ниже безопасной себестоимости.")
    add_table(doc, ["Показатель", "Рабочее правило"], [
        ["Выручка", "Подписка организации + амортизированная плата за внедрение + платная поддержка или кастомизация."],
        ["Переменные расходы", "Запросы к модели + STT и TTS + трафик и backend + поддержка, зависящая от числа пользователей."],
        ["Целевая маржа", "Не ниже 70 процентов после завершения пилота."],
        ["Минимальная цена", "Если V - переменная стоимость пользователя, переменная часть цены должна быть не меньше V / 0,30."],
        ["Ограничения", "В пакете нужен лимит включённых вопросов и голосовых минут; дополнительный объём оформляется как add-on или мягкое ограничение."],
    ], [4.2, 12.8])
    add_body(doc, "Перед любым коммерческим предложением нужно подставить актуальные тарифы выбранных AI и Speech-провайдеров на дату расчёта. Публичные устаревшие цены не использовать.", lead="Дисциплина расчёта. ")

    add_heading(doc, "6  Первые 90 дней")
    add_table(doc, ["Этап", "Период", "Действия", "Результат"], [
        ["Подтверждение боли", "Недели 1-3", "10-15 интервью с пользователями, преподавателями, библиотекарями и сотрудниками поддержки. Проверка 2-3 безопасных сценариев.", "Список реальных барьеров и критерии пилота."],
        ["Контролируемый пилот", "Недели 4-8", "Один партнёр, 20-50 участников, учебные страницы и короткая вводная о границах помощника.", "Данные о завершении задач, понимании подсказок и сбоях."],
        ["Решение о масштабе", "Недели 9-12", "Сравнение до и после, улучшение трёх самых частых сценариев, коммерческое предложение партнёру.", "Решение о платном продолжении или корректировке продукта."],
    ], [3.0, 2.3, 7.1, 4.6])

    add_heading(doc, "Метрики пилота", level=2)
    add_table(doc, ["Метрика", "Что показывает"], [
        ["Activation rate", "Понятность первого опыта: сколько установивших расширение прошли первое обучение или задали вопрос."],
        ["Task completion", "Практическую пользу: долю завершённых безопасных сценариев среди начавших."],
        ["Help deflection", "Потенциальное снижение типовых вопросов консультанту."],
        ["Repeat success", "Обучающий эффект: может ли человек повторить задачу без подсказки."],
        ["KK adoption и Voice success", "Востребованность казахского интерфейса и надёжность голосового пути."],
        ["Safety precision review", "Не создаёт ли предупреждение о риске лишнюю тревогу."],
    ], [4.3, 12.7])
    add_body(doc, "Собирать следует только агрегированные продуктовые события. Полный текст страницы, документы, пароли, записи голоса и содержимое личных аккаунтов не должны входить в аналитику. Для production-аналитики требуется отдельное понятное согласие.")
    doc.add_page_break()

    add_heading(doc, "7  Риски и принципы доверия")
    add_table(doc, ["Риск", "Как управлять"], [
        ["Неверная подсказка модели", "Один безопасный следующий шаг, строгая схема ответа, запрет автоматических действий, просьба открыть меню при недостатке данных."],
        ["Чувствительный контент", "Локальная фильтрация, минимальный контекст, запрет чтения документов и полей, понятная политика данных."],
        ["Ложное чувство безопасности", "Формулировка о возможной угрозе, отсутствие зелёного статуса для чистой страницы, обучение пользователя."],
        ["Рост облачных расходов", "Лимиты, cache проверенных ответов, мониторинг затрат и пакеты по объёму."],
        ["Динамические интерфейсы", "Честное ограничение, повторный анализ после открытия меню и развитие покрытия по данным пилота."],
        ["Качество казахского", "Проверка носителями, библиотека терминов и тесты ключевых сценариев."],
    ], [4.3, 12.7])

    add_heading(doc, "8  Сообщение для жюри")
    add_body(doc, "ИИ Наставник помогает человеку не просто найти ответ, а уверенно сделать следующий шаг в интерфейсе. Он работает с ограниченными описаниями видимых элементов, объясняет их на казахском или русском и ничего не нажимает вместо пользователя. Поэтому решение подходит начинающим пользователям, пожилым, школьникам и организациям цифровой грамотности. Мы начинаем с B2B2C-пилотов в школах, библиотеках и учебных центрах, измеряем завершение реальных задач и снижение однотипных обращений за помощью. Доход формируют организационная подписка и внедрение, а безопасность и приватность остаются архитектурными ограничениями, а не дополнительной платной опцией.")

    add_heading(doc, "Следующий шаг для команды", level=2)
    add_body(doc, "Выбрать одну стартовую вертикаль - образовательную организацию или библиотеку с программой цифровой грамотности - и согласовать пилот с измеримой метрикой успеха. После пилота обновить этот документ фактическими данными: размером группы, завершением сценариев, стоимостью активного пользователя, отзывами о KK и RU и решением партнёра о продлении.", lead="Решение. ")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "Бизнес модель ИИ Наставника"
    doc.core_properties.subject = "Пилотная и коммерческая модель проекта"
    doc.core_properties.author = "ИИ Наставник"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
