import csv

from openpyxl import load_workbook

from backend.demo import create_demo
from backend.engine import analyze


def test_xlsx_keeps_all_ids_cyrillic_and_numeric_values(tmp_path):
    run = analyze(create_demo(tmp_path / 'input'), tmp_path / 'out', synthetic=True)
    book = load_workbook(run.output_dir / 'results.xlsx', data_only=False)
    assert book.sheetnames == ['Приоритеты', 'Все счета', 'Сообщества', 'О выгрузке']
    sheet = book['Все счета']
    assert sheet.max_row == len(run.nodes) + 1
    assert sheet.freeze_panes == 'A2'
    assert sheet.auto_filter.ref == f'A1:F{len(run.nodes)+1}'
    for row, node in zip(sheet.iter_rows(min_row=2), run.nodes):
        assert row[0].data_type == 's' and row[0].value == node['gid']
        assert row[0].number_format == '@'
        assert row[2].data_type == 'n' and row[2].value == node['role_score']
        assert row[4].data_type == 'n' and row[4].value == node['priority_score']
        assert row[5].value == node['evidence']
    assert book['Приоритеты'].max_row == min(100, len(run.nodes)) + 1
    for row, node in zip(book['Приоритеты'].iter_rows(min_row=2), run.nodes):
        assert row[1].data_type == 's' and row[1].value == node['gid']
    for row, cluster in zip(book['Сообщества'].iter_rows(min_row=2), run.clusters):
        assert row[3].data_type == 'n' and row[3].value == cluster['sum_kzt_internal']
        assert row[4].data_type == 's' and row[4].value == cluster['top_gids']
        assert row[5].value == cluster['hypothesis']
    assert all(cell.data_type != 'f' for sheet in book for row in sheet for cell in row)
    book.close()


def test_csv_has_utf8_bom_and_preserves_schema_and_text(tmp_path):
    run = analyze(create_demo(tmp_path / 'input'), tmp_path / 'out')
    for filename in ['nodes_roles.csv', 'clusters.csv', 'top_nodes.csv']:
        assert (run.output_dir / filename).read_bytes().startswith(b'\xef\xbb\xbf')
    with (run.output_dir / 'nodes_roles.csv').open(encoding='utf-8-sig', newline='') as file:
        records = list(csv.DictReader(file))
    assert len(records) == len(run.nodes)
    assert records[0]['gid'] == run.nodes[0]['gid']
    assert records[0]['evidence'] == run.nodes[0]['evidence']


def test_formula_like_metadata_remains_literal_text(tmp_path):
    run = analyze(create_demo(tmp_path / 'input'), tmp_path / 'out', label='=HYPERLINK("https://example.invalid","Кириллица")')
    book = load_workbook(run.output_dir / 'results.xlsx', data_only=False)
    cell = book['О выгрузке']['B2']
    assert cell.data_type == 's' and cell.value == run.summary['label']
    book.close()
